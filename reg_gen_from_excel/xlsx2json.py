#!/usr/bin/env python3
# FILE NAME      : xlsx2json.py
# DEPARTMENT     : Digital Verification Department
# AUTHOR         : zhouxk
# AUTHOR'S EMAIL : zhouxk@chipsea.com
# ABSTRACT       : Excel-to-JSON conversion utilities. Handles merged cells,
#                  blank row filtering, column validation, and nested JSON
#                  construction grouped by register name.
#
#-------------------------------------------------------------------------

import pandas as pd
import openpyxl
import os
import time
import json
import sys

# ---------------------------------------------------------------------------
# Merged cell handling
# ---------------------------------------------------------------------------

def unmerge_cells(worksheet):
    """Unmerge all merged cells in a worksheet, filling each cell in the
    merged range with the value from the top-left cell."""
    list_ranges = list(worksheet.merged_cells.ranges)
    for range_obj in list_ranges:
        cell_start = range_obj.start_cell
        worksheet.unmerge_cells(range_string=range_obj.coord)
        for row_index, col_index in range_obj.cells:
            cell = worksheet.cell(row=row_index, column=col_index)
            cell.value = cell_start.value

def unmerge_cell(filename, sheet_name):
    """Load an Excel file, unmerge all cells in the target sheet, save to a
    temporary file, and return the temp file path. Returns None on error."""
    try:
        if not os.path.exists(filename):
            raise FileNotFoundError(f"{filename} not found")
        wb = openpyxl.load_workbook(filename)
        sheet = wb[sheet_name]
        unmerge_cells(sheet)
        temp_filename = f'unmerge_cell_temp_{int(time.time())}.xlsx'
        wb.save(temp_filename)
        wb.close()
        return temp_filename
    except Exception as e:
        print(f"Excel error: {e}")
        return None

# ---------------------------------------------------------------------------
# Blank row filtering
# ---------------------------------------------------------------------------

def filter_blank_rows(df, key_columns):
    """Filter out rows where ALL key columns are empty/NaN.
    Uses OR logic: a row is kept if ANY key column has a value."""
    before_count = len(df)
    present_cols = [col for col in key_columns if col in df.columns]
    if not present_cols:
        return df
    mask = pd.Series([False] * len(df), index=df.index)
    for col in present_cols:
        col_valid = df[col].notna() & (df[col].astype(str).str.strip() != '')
        mask = mask | col_valid
    df_filtered = df[mask]
    skipped = before_count - len(df_filtered)
    if skipped > 0:
        print(f"\033[1;33m[Blank Row Filter]\033[0m Skipped {skipped} "
              f"blank/incomplete row(s) (key columns: {present_cols})")
    return df_filtered

# ---------------------------------------------------------------------------
# Column validation
# ---------------------------------------------------------------------------

def validate_excel_columns(df, expected_columns):
    """Compare actual Excel column names against expected names from YAML.
    Strips whitespace/newlines from column names. Reports mismatches."""
    excel_cols = [str(c).strip().replace("\n", "").replace("\r", "")
                  for c in df.columns]
    df.columns = excel_cols

    missing = []
    mismatched = []
    for exp in expected_columns:
        exp_stripped = exp.strip()
        if exp_stripped in excel_cols:
            continue
        found_partial = False
        for ec in excel_cols:
            if (exp_stripped.lower() in ec.lower()
                    or ec.lower() in exp_stripped.lower()):
                mismatched.append((exp_stripped, ec))
                found_partial = True
                break
        if not found_partial:
            missing.append(exp_stripped)

    if mismatched or missing:
        print("\033[1;33m[Column Mismatch Warning]\033[0m")
        print(f"  Excel columns  : {excel_cols}")
        print(f"  Expected columns: {expected_columns}")
        if mismatched:
            print("  \033[1;31mSimilar but not matching:\033[0m")
            for exp, actual in mismatched:
                print(f"    Expected '{exp}' -> Found '{actual}'")
        if missing:
            print("  \033[1;31mMissing columns:\033[0m")
            for m in missing:
                print(f"    '{m}'")
        print("  Please ensure Excel header row matches YAML config "
              "exactly (watch for spaces/newlines).")
    return df

# ---------------------------------------------------------------------------
# Main Excel-to-JSON conversion
# ---------------------------------------------------------------------------

def excel_to_nested_json(excel_file_path, sheet_name, skip_rows=0,
                         usecols=None, nrows=None,
                         expected_columns=None, key_columns=None):
    """Read an Excel sheet and convert to a list of row dicts.
    Handles merged cells, column validation, and blank row filtering."""
    # Step 1: Unmerge cells and save to temp file
    temp_file = unmerge_cell(excel_file_path, sheet_name)
    if temp_file is None:
        raise ValueError(
            f"Failed to unmerge cells in "
            f"{excel_file_path}/{sheet_name}")

    # Step 2: Read Excel into DataFrame
    try:
        df = pd.read_excel(
            temp_file, sheet_name=sheet_name, skiprows=skip_rows,
            usecols=usecols, nrows=nrows, engine='openpyxl')
    except Exception as e:
        raise ValueError(f"Failed to read Excel file: {str(e)}")
    finally:
        # [FIX #17] Always clean up temp file, even on error
        if temp_file and os.path.exists(temp_file):
            os.remove(temp_file)

    # Step 3: Validate column names against YAML config
    if expected_columns:
        df = validate_excel_columns(df, expected_columns)

    # Step 4: Filter blank rows
    if key_columns:
        df = filter_blank_rows(df, key_columns)

    # Step 5: Forward-fill to handle remaining merged cells
    df = df.ffill()
    nested_data = df.to_dict(orient='records')
    return nested_data

# ---------------------------------------------------------------------------
# JSON merging by primary key
# ---------------------------------------------------------------------------

def merge_json_by_key(raw_json, primary_key, common_columns, group_name):
    """Merge flat row list into nested structure grouped by primary_key.
    Each group contains common columns + a list of field_group entries.
    [FIX #10] Raises KeyError instead of calling sys.exit()."""
    merged = {}
    has_find = False

    for item in raw_json:
        item_name = item.get(primary_key)
        if not item_name:
            continue
        has_find = True
        if item_name not in merged:
            # [FIX #10] Use explicit check instead of sys.exit in dict comprehension
            col_values = {}
            for col in common_columns:
                val = item.get(col)
                if val is None:
                    raise KeyError(
                        f"\033[1;31mColumn [{col}] not found in Excel. "
                        f"Please match YAML and Excel headers.\033[0m")
                col_values[col] = val
            merged[item_name] = {
                primary_key: item_name,
                **col_values,
                group_name: []
            }
        # Collect non-common columns as field entries
        non_common = {k: v for k, v in item.items()
                      if k not in [primary_key] + common_columns
                      and v is not None}
        if non_common:
            merged[item_name][group_name].append(non_common)

    if not has_find:
        raise ValueError(
            f"\033[1;31mCannot find '{primary_key}' in Excel data\033[0m")
    return list(merged.values())


# ---------------------------------------------------------------------------
# Standalone test entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    config = {
        "excel_file_path": "./Airbag_register_map.xlsx",
        "sheet_name": "DEPLOYMENT_DRIVER01",
        "skip_rows": 2,
        "usecols": "A,C:H,J:L",
        "nrows": None
    }
    json_obj = excel_to_nested_json(**config)
    primary_key = "Register Name"
    common_columns = ["Address Offset", "Reset_Value", "Read", "Write", "Reset"]
    group_name = "field_group"
    result = merge_json_by_key(json_obj, primary_key, common_columns, group_name)
    print(json.dumps(result, indent=2, ensure_ascii=False))
